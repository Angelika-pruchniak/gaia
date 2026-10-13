# Generuje Dla_prowadzacej/Karta_analizy_konsultacji.xlsx (analiza toku rozumowania + raport dla klienta)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
import os

NAVY='1E2A4A'; ORANGE='F8B68C'; PEACH='FDDCC6'; PINK='F2BFC5'; BLUSH='FBE6E7'; CREAM='FFF8F3'; GREY='6B6F80'
F=lambda c: PatternFill('solid', start_color=c, end_color=c)
thin=Side(style='thin', color='E3E5EC'); B=Border(left=thin,right=thin,top=thin,bottom=thin)
WR=Alignment(wrap_text=True,vertical='top')
wb=Workbook()

def title(ws,t,sub,n=6):
    ws.sheet_view.showGridLines=False; ws.column_dimensions['A'].width=2
    ws['B2']=t; ws['B2'].font=Font(size=18,bold=True,color=NAVY)
    ws['B3']=sub; ws['B3'].font=Font(size=10,italic=True,color=GREY)
    for c in range(2,2+n): ws.cell(row=1,column=c).fill=F(ORANGE if c%2 else PINK)
    ws.row_dimensions[1].height=8
def inp(c,v=None,fmt=None,h=None):
    if v is not None: c.value=v
    c.fill=F(PEACH); c.border=B; c.alignment=WR; c.font=Font(color=NAVY)
    if fmt: c.number_format=fmt
def lab(c,t,bold=True,fill=None):
    c.value=t; c.font=Font(bold=bold,color=NAVY); c.alignment=WR; c.border=B
    if fill: c.fill=F(fill)
def sec(ws,r,t,cols='BCDE'):
    ws[f'B{r}']=t
    for col in cols: ws[f'{col}{r}'].fill=F(NAVY); ws[f'{col}{r}'].font=Font(bold=True,color='FFFFFF')

# ---------- Instrukcja ----------
ws=wb.active; ws.title='Instrukcja'
title(ws,'Przegląd Decyzji · karta analizy','Narzędzie prowadzącej. Jeden plik = jeden klient. Zapisz kopię jako „Nazwa klienta_data.xlsx”.',4)
ws.column_dimensions['B'].width=22; ws.column_dimensions['C'].width=90
steps=[('1. Klient','Dane z karty przygotowania i kalkulatora klienta (arkusz Podsumowanie).'),
('2. Analiza','Przed sesją: dla 3 decyzji oceń 6 ogniw (1–5), wpisz ukryte założenia, sygnały ostrzegawcze, pytania i hipotezę. Arkusz wskaże najsłabsze ogniwo i kolejność omawiania.'),
('3. Sesja','Pracuj według scenariusza 120 min (Pakiet prowadzącej, rozdz. 6). Po sesji uzupełnij werdykty, zmiany i wzorzec decyzyjny.'),
('4. Raport','Arkusz Raport składa się sam. Uzupełnij pola brzoskwiniowe (podsumowanie, rekomendacje, plan 30 dni), potem Plik → Drukuj → PDF i wyślij klientowi w 3 dni robocze.'),
('Skala ogniw','1–2 słabe · 3 wystarczające · 4–5 mocne (kryteria: Pakiet prowadzącej, rozdz. 5).'),
('Przykład','Plik zawiera przykładowe dane fikcyjnego klienta „Studio Lumen”. Usuń je przed pierwszym użyciem.')]
for i,(a,b) in enumerate(steps):
    r=5+i; lab(ws[f'B{r}'],a); ws[f'C{r}']=b; ws[f'C{r}'].alignment=WR; ws[f'C{r}'].border=B; ws.row_dimensions[r].height=34
    if i%2==0: ws[f'B{r}'].fill=F(BLUSH); ws[f'C{r}'].fill=F(BLUSH)

# ---------- Klient ----------
ws=wb.create_sheet('Klient')
title(ws,'Klient','Dane z karty przygotowania do konsultacji',3)
ws.column_dimensions['B'].width=42; ws.column_dimensions['C'].width=70
fields=[('Klient (osoba)','Anna Nowak (przykład)'),('Firma','Studio Lumen (fikcyjne)'),('Branża / liczba osób','Studio graficzno-marketingowe, 6 osób'),
('Data sesji','2026-11-20'),('Przychód / marża netto (12 mies.)','1,2 mln zł / 8%'),('Wynik audytu / najsłabszy obszar','31/60 · Oferta i ceny'),
('Cel na 12 miesięcy','Marża netto 18%, 40% przychodu z abonamentów'),('Czego oczekuje po konsultacji','Pewności, że podwyżka nie zniszczy relacji z największymi klientami'),
('Decyzja z najmniejszą pewnością','Podwyżka 10% dla wszystkich obecnych klientów'),('Inni decydenci','Wspólnik (udziały 30%), nie uczestniczył w programie')]
for i,(a,b) in enumerate(fields):
    r=5+i; lab(ws[f'B{r}'],a); inp(ws[f'C{r}'],b); ws.row_dimensions[r].height=30

# ---------- Analiza ----------
ws=wb.create_sheet('Analiza')
title(ws,'Analiza toku rozumowania','Wiersze = elementy analizy, kolumny = decyzje. Brzoskwiniowe pola wypełniasz.',4)
ws.column_dimensions['B'].width=34
for c in 'CDE': ws.column_dimensions[c].width=42
ws['B5']=''; 
for j,c in enumerate('CDE'):
    ws[f'{c}5']=f'Decyzja {j+1}'; ws[f'{c}5'].font=Font(bold=True,color='FFFFFF'); ws[f'{c}5'].fill=F(NAVY); ws[f'{c}5'].alignment=Alignment(horizontal='center')
ws['B5'].fill=F(NAVY)
ex={'C':['Gdzie grać: segment kliniki i gabinety, pakiet flagowy','Ceny: podwyżka 10% dla wszystkich od 1.03','Priorytety: abonament + standard sprzedaży'],
}
dec=['Gdzie grać: kliniki i gabinety + pakiet flagowy','Ceny: podwyżka 10% dla wszystkich od 1.03','3 priorytety: abonament, cennik, standard sprzedaży']
r=6; lab(ws[f'B{r}'],'Decyzja (z karty klienta)')
for j,c in enumerate('CDE'): inp(ws[f'{c}{r}'],dec[j])
ws.row_dimensions[r].height=34
sec(ws,8,'OCENA 6 OGNIW (1–5)')
links=['Rama','Opcje','Informacje','Kryteria','Rozumowanie','Zobowiązanie']
scores=[[4,3,3],[2,1,3],[4,3,2],[3,2,3],[3,2,3],[3,4,2]]
dv=DataValidation(type='whole',operator='between',formula1='1',formula2='5',allow_blank=True); ws.add_data_validation(dv)
for i,l in enumerate(links):
    rr=9+i; lab(ws[f'B{rr}'],l)
    for j,c in enumerate('CDE'): inp(ws[f'{c}{rr}'],scores[i][j]); ws[f'{c}{rr}'].alignment=Alignment(horizontal='center'); dv.add(f'{c}{rr}')
ws.conditional_formatting.add('C9:E14',CellIsRule(operator='lessThanOrEqual',formula=['2'],fill=F(PINK),font=Font(bold=True,color='9C2B3B')))
ws.conditional_formatting.add('C9:E14',CellIsRule(operator='greaterThanOrEqual',formula=['4'],fill=F(ORANGE)))
lab(ws['B15'],'Średnia',fill=CREAM); lab(ws['B16'],'Najsłabsze ogniwo',fill=CREAM); lab(ws['B17'],'Kolejność omawiania (1 = najpierw)',fill=CREAM)
for c in 'CDE':
    ws[f'{c}15']=f'=IFERROR(ROUND(AVERAGE({c}9:{c}14),1),"")'
    ws[f'{c}16']=f'=IFERROR(INDEX($B$9:$B$14,MATCH(MIN({c}9:{c}14),{c}9:{c}14,0))&" ("&MIN({c}9:{c}14)&")","")'
    ws[f'{c}17']=f'=IFERROR(RANK({c}15,$C$15:$E$15,1),"")'
    for rr in (15,16,17): ws[f'{c}{rr}'].border=B; ws[f'{c}{rr}'].font=Font(bold=True,color=NAVY); ws[f'{c}{rr}'].alignment=Alignment(horizontal='center'); ws[f'{c}{rr}'].fill=F(CREAM)
sec(ws,19,'NOTATKI DO OGNIW (dlaczego taka ocena)')
notes=[['Dane z kalkulatora potwierdzają: najwyższy zysk/h w medycznych','Jedna opcja: +10% dla wszystkich; brak wariantu dla Gwiazd','Lista pomysłów obszerna, wybór sensowny'],
['Tylko 2 rozmowy z klientami z segmentu','Brak danych o elastyczności: klient A (18% przychodu) bez rozmowy','Brak oszacowania czasu wdrożenia'],
['Pre-mortem konkretny','Pre-mortem: „nic nie pójdzie źle”','Trzy priorytety wymagają pracy właścicielki']]
for k,lbl in enumerate(['Mocne strony','Luki w informacjach','Zobowiązanie / wdrożenie']):
    rr=20+k; lab(ws[f'B{rr}'],lbl)
    for j,c in enumerate('CDE'): inp(ws[f'{c}{rr}'],notes[k][j])
    ws.row_dimensions[rr].height=44
sec(ws,24,'UKRYTE ZAŁOŻENIA (klient ich nie zapisał)')
ass=[['Gabinety mają budżet na pakiet 16 900 zł','Klient A nie ma alternatywy w tej cenie','Właścicielka znajdzie 6 h tygodniowo'],
['Partner (biura rachunkowe) poleca aktywnie','Wszyscy klienci reagują na podwyżkę tak samo','PM przejmie rozmowy bez szkolenia'],
['','Wspólnik zaakceptuje ryzyko utraty klienta A','']]
for k in range(3):
    rr=25+k; lab(ws[f'B{rr}'],f'Założenie {k+1}')
    for j,c in enumerate('CDE'): inp(ws[f'{c}{rr}'],ass[k][j])
    ws.row_dimensions[rr].height=34
sec(ws,29,'SYGNAŁY OSTRZEGAWCZE (wpisz X)')
flags=['Jedna opcja jako jedyna możliwa','„Wszyscy w branży tak robią”','Cena względem konkurencji / dawnego cennika','Decyzja chroni przed stratą, nie buduje zysku','Cel bez liczby lub terminu','Pewność 90%+ przy słabych danych','Każda decyzja wymaga pracy właściciela','Pre-mortem pusty','Trzymanie nierentownego klienta z sentymentu','Wdrożenie „przy okazji”']
fx={1:'D',3:'D',5:'D',6:'E',7:'D',9:'E'}
for k,fl in enumerate(flags):
    rr=30+k; lab(ws[f'B{rr}'],fl,bold=False)
    for c in 'CDE': inp(ws[f'{c}{rr}'],'X' if fx.get(k)==c or (k==0 and c=='D') else None); ws[f'{c}{rr}'].alignment=Alignment(horizontal='center')
lab(ws['B40'],'Liczba sygnałów',fill=CREAM)
for c in 'CDE': ws[f'{c}40']=f'=COUNTIF({c}30:{c}39,"X")'; ws[f'{c}40'].font=Font(bold=True,color=NAVY); ws[f'{c}40'].border=B; ws[f'{c}40'].alignment=Alignment(horizontal='center'); ws[f'{c}40'].fill=F(CREAM)
ws.conditional_formatting.add('C40:E40',CellIsRule(operator='greaterThanOrEqual',formula=['2'],fill=F(PINK)))
sec(ws,42,'PYTANIA NA SESJĘ')
qs=[['Co musiałoby być prawdą, żeby gabinety kupowały pakiet za 16 900 zł?','Co by Pani zrobiła, gdyby podwyżka dla klienta A była niemożliwa?','Co Pani przestanie robić, żeby znaleźć 6 h tygodniowo?'],
['Który segment odrzuciła Pani najszybciej i dlaczego?','Czego najbardziej się Pani obawia w rozmowie z klientem A?','Kto poza Panią może prowadzić rozmowy o abonamencie?'],
['Jaki fakt zmieniłby Pani zdanie co do segmentu?','Co doradziłaby Pani znajomej właścicielce w tej sytuacji?','Po czym pozna Pani w 4. tygodniu, że plan nie działa?']]
for k in range(3):
    rr=43+k; lab(ws[f'B{rr}'],f'Pytanie {k+1}')
    for j,c in enumerate('CDE'): inp(ws[f'{c}{rr}'],qs[k][j])
    ws.row_dimensions[rr].height=44
sec(ws,47,'HIPOTEZA PRZED SESJĄ')
lab(ws['B48'],'Moja hipoteza (zostaw / popraw / wstrzymaj + dlaczego)')
hyp=['Zostaw. Mocne dane; dodać test z 5 gabinetami','Popraw. Zróżnicować podwyżkę: Gwiazdy +5% z wartością dodaną, „Do pożegnania” +15%, klient A rozmowa osobista','Popraw. Odciążyć właścicielkę: standard sprzedaży jako priorytet 1']
for j,c in enumerate('CDE'): inp(ws[f'{c}48'],hyp[j])
ws.row_dimensions[48].height=60
sec(ws,50,'PO SESJI')
vd=DataValidation(type='list',formula1='"Zostaw,Popraw,Wstrzymaj"',allow_blank=True); ws.add_data_validation(vd)
lab(ws['B51'],'Werdykt'); lab(ws['B52'],'Co konkretnie zmieniamy'); lab(ws['B53'],'Pierwszy krok, kto, kiedy')
post=[['Zostaw','Popraw','Popraw'],['Dodać test: 5 rozmów z gabinetami do 15.12','Podwyżka zróżnicowana; klient A: rozmowa właścicielki do 10.01, wariant zakresu','Kolejność: 1) standard sprzedaży 2) cennik 3) abonament'],['Właścicielka: lista 20 gabinetów do 5.12','Właścicielka + wspólnik: decyzja o kliencie A do 20.12','PM: skrypty i CRM do 31.12']]
for k,rr in enumerate((51,52,53)):
    for j,c in enumerate('CDE'): inp(ws[f'{c}{rr}'],post[k][j])
    ws.row_dimensions[rr].height=40 if rr>51 else 20
for c in 'CDE': vd.add(f'{c}51')
ws.conditional_formatting.add('C51:E51',FormulaRule(formula=['C51="Zostaw"'],fill=F(ORANGE)))
ws.conditional_formatting.add('C51:E51',FormulaRule(formula=['C51="Popraw"'],fill=F(PEACH)))
ws.conditional_formatting.add('C51:E51',FormulaRule(formula=['C51="Wstrzymaj"'],fill=F(PINK)))
lab(ws['B55'],'Wzorzec decyzyjny klienta'); ws.merge_cells('C55:E55'); inp(ws['C55'],'Przy decyzjach o cenach i ludziach wygrywa obawa przed konfliktem: zawęża opcje do jednej „bezpiecznej” i pomija rozmowę z najważniejszym interesariuszem.'); ws.row_dimensions[55].height=40
lab(ws['B56'],'Nawyk na przyszłość'); ws.merge_cells('C56:E56'); inp(ws['C56'],'Przed każdą decyzją powyżej 10 tys. zł: 3 opcje (w tym status quo) + pre-mortem + rozmowa z osobą, której decyzja dotyczy najbardziej.'); ws.row_dimensions[56].height=34
ws.freeze_panes='C6'

# ---------- Raport ----------
ws=wb.create_sheet('Raport')
title(ws,'Raport z konsultacji Przegląd Decyzji','guliwera4.pl · poufne',5)
for c,w in zip('BCDEF',[30,22,16,22,40]): ws.column_dimensions[c].width=w
ws['B4']='=Klient!C6&" · "&Klient!C5'; ws['B4'].font=Font(bold=True,size=12,color=NAVY)
ws['F4']='="Data sesji: "&Klient!C8'; ws['F4'].font=Font(color=GREY); ws['F4'].alignment=Alignment(horizontal='right')
def rsec(r,t):
    ws[f'B{r}']=t
    for c in 'BCDEF': ws[f'{c}{r}'].fill=F(NAVY); ws[f'{c}{r}'].font=Font(bold=True,color='FFFFFF')
rsec(6,'1. PODSUMOWANIE')
ws.merge_cells('B7:F7'); inp(ws['B7'],'Firma ma mocny fundament: wybrany segment medyczny daje najwyższy zysk na godzinę, a dane z kalkulatora go potwierdzają. Najsłabszym miejscem planu jest decyzja o cenach: oparta na jednej opcji i na obawie przed utratą klienta A, a nie na jego realnej sytuacji. Plan wymaga też odciążenia właścicielki, żeby był wykonalny.'); ws.row_dimensions[7].height=70
rsec(9,'2. MAPA DECYZJI')
hdr=['Decyzja','Średnia ocena (1–5)','Najsłabsze ogniwo','Werdykt','Co zmieniamy']
for j,h in enumerate(hdr):
    c=ws.cell(row=10,column=2+j,value=h); c.font=Font(bold=True,color=NAVY); c.fill=F(ORANGE); c.border=B; c.alignment=Alignment(wrap_text=True,vertical='center')
for k,c in enumerate('CDE'):
    rr=11+k
    ws[f'B{rr}']=f'=Analiza!{c}6'; ws[f'C{rr}']=f'=Analiza!{c}15'; ws[f'D{rr}']=f'=Analiza!{c}16'; ws[f'E{rr}']=f'=Analiza!{c}51'; ws[f'F{rr}']=f'=Analiza!{c}52'
    for col in 'BCDEF': ws[f'{col}{rr}'].border=B; ws[f'{col}{rr}'].alignment=WR; ws[f'{col}{rr}'].font=Font(color=NAVY)
    ws[f'E{rr}'].font=Font(bold=True,color=NAVY); ws.row_dimensions[rr].height=48
ws.conditional_formatting.add('E11:E13',FormulaRule(formula=['E11="Zostaw"'],fill=F(ORANGE)))
ws.conditional_formatting.add('E11:E13',FormulaRule(formula=['E11="Popraw"'],fill=F(PEACH)))
ws.conditional_formatting.add('E11:E13',FormulaRule(formula=['E11="Wstrzymaj"'],fill=F(PINK)))
rsec(15,'3. OCENA 6 OGNIW')
for j,h in enumerate(['Ogniwo','Decyzja 1','Decyzja 2','Decyzja 3']):
    c=ws.cell(row=16,column=2+j,value=h); c.font=Font(bold=True,color=NAVY); c.fill=F(ORANGE); c.border=B; c.alignment=Alignment(horizontal='center')
for i,l in enumerate(links):
    rr=17+i; ws[f'B{rr}']=l; ws[f'B{rr}'].border=B; ws[f'B{rr}'].font=Font(bold=True,color=NAVY)
    for j,c in enumerate('CDE'):
        cell=ws.cell(row=rr,column=3+j,value=f'=Analiza!{c}{9+i}'); cell.border=B; cell.alignment=Alignment(horizontal='center')
ws.conditional_formatting.add('C17:E22',CellIsRule(operator='lessThanOrEqual',formula=['2'],fill=F(PINK),font=Font(bold=True,color='9C2B3B')))
ws.conditional_formatting.add('C17:E22',CellIsRule(operator='greaterThanOrEqual',formula=['4'],fill=F(ORANGE)))
ws['F17']='Skala: 1–2 słabe · 3 wystarczające · 4–5 mocne. Decyzja jest tak mocna, jak jej najsłabsze ogniwo.'; ws['F17'].font=Font(size=9,italic=True,color=GREY); ws['F17'].alignment=WR; ws.merge_cells('F17:F22')
rsec(24,'4. ZAŁOŻENIA DO SPRAWDZENIA')
for j,h in enumerate(['Założenie','Jak sprawdzić','','Do kiedy','Kto']):
    c=ws.cell(row=25,column=2+j,value=h); c.font=Font(bold=True,color=NAVY); c.fill=F(ORANGE); c.border=B
ws.merge_cells('C25:D25')
ex_as=[('=Analiza!D26','Rozmowy z 3 klientami z ćwiartki Gwiazd przed ogłoszeniem podwyżki','15.12','Właścicielka'),('=Analiza!C25','5 rozmów z gabinetami (karta testu)','15.12','Właścicielka'),('=Analiza!E26','Szkolenie PM + 2 rozmowy pod obserwacją','31.12','PM')]
for k,(a,b,d,e) in enumerate(ex_as):
    rr=26+k; ws[f'B{rr}']=a; ws[f'B{rr}'].border=B; ws[f'B{rr}'].alignment=WR
    ws.merge_cells(f'C{rr}:D{rr}'); inp(ws[f'C{rr}'],b); inp(ws[f'E{rr}'],d); inp(ws[f'F{rr}'],e); ws.row_dimensions[rr].height=34
rsec(30,'5. TWÓJ WZORZEC DECYZYJNY')
ws.merge_cells('B31:F31'); ws['B31']='=Analiza!C55'; ws['B31'].alignment=WR; ws.row_dimensions[31].height=40
ws.merge_cells('B32:F32'); ws['B32']='="Nawyk na przyszłość: "&Analiza!C56'; ws['B32'].alignment=WR; ws['B32'].font=Font(bold=True,color=NAVY); ws.row_dimensions[32].height=34
rsec(34,'6. TRZY REKOMENDACJE')
recs=['Zróżnicuj podwyżkę według ćwiartek macierzy i przeprowadź osobistą rozmowę z klientem A, zanim wyślesz list.','Uczyń standard sprzedaży (skrypty, CRM, PM) priorytetem nr 1. Bez niego pozostałe dwa priorytety zależą wyłącznie od Ciebie.','Potwierdź segment medyczny testem 5 rozmów, zanim zmienisz stronę i komunikację.']
for k in range(3):
    rr=35+k; ws[f'B{rr}']=f'{k+1}.'; ws[f'B{rr}'].font=Font(bold=True,color=NAVY); ws[f'B{rr}'].alignment=Alignment(horizontal='right',vertical='top')
    ws.merge_cells(f'C{rr}:F{rr}'); inp(ws[f'C{rr}'],recs[k]); ws.row_dimensions[rr].height=34
rsec(39,'7. PLAN NA 30 DNI')
for j,h in enumerate(['Tydzień','Działanie','','','Kto']):
    c=ws.cell(row=40,column=2+j,value=h); c.font=Font(bold=True,color=NAVY); c.fill=F(ORANGE); c.border=B
ws.merge_cells('C40:E40')
plan=[('Tydzień 1','Lista 20 gabinetów; rozmowy z 3 Gwiazdami o zmianach','Właścicielka'),('Tydzień 2','Rozmowa z klientem A; decyzja ze wspólnikiem','Właścicielka + wspólnik'),('Tydzień 3','Skrypty i CRM; PM prowadzi 2 rozmowy','PM'),('Tydzień 4','Wysyłka listów o zmianie cen; przegląd wyników testu','Właścicielka')]
for k,(a,b,c_) in enumerate(plan):
    rr=41+k; inp(ws[f'B{rr}'],a); ws.merge_cells(f'C{rr}:E{rr}'); inp(ws[f'C{rr}'],b); inp(ws[f'F{rr}'],c_); ws.row_dimensions[rr].height=30
rsec(46,'8. KOLEJNY KONTAKT')
ws.merge_cells('B47:F47'); inp(ws['B47'],'Pytania mailowe do 14 dni po sesji. Przegląd po 30 dniach (30 min): [data i godzina].'); ws.row_dimensions[47].height=30
ws.print_area='B2:F47'; ws.page_setup.orientation='portrait'; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
ws.sheet_properties.pageSetUpPr.fitToPage=True; ws.print_options.horizontalCentered=True
ws.page_margins.left=ws.page_margins.right=0.4

for s in wb.worksheets: s.sheet_properties.tabColor = NAVY if s.title in ('Instrukcja','Raport') else ORANGE
out=os.path.join(os.path.dirname(__file__),'..','Dla_prowadzacej','Karta_analizy_konsultacji.xlsx')
wb.save(out); print('saved',out)
